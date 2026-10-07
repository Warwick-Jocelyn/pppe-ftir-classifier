"""Build Dataset/Processed: a self-contained, human-checkable PP/PE dataset with explicit train / val / test splits.

Everything is derived from the raw files in Dataset/ (which are left untouched):
  * GT sheet   `FTIR Classification_GT_Labels.xlsx`  (label = LAST column "Varified Polymer ID")
  * Colour     `Colorimeter_data.xlsx`               (CIELAB L*, a*, b*, D65 / 10 deg)
  * Spectra    OPUS exports `<sample_id>.0.dpt` (+ the native `<sample_id>.0`)
Rows are joined by SAMPLE ID (never by row order) and every record keeps the Excel row numbers it came from,
so it can be traced back by hand. `verify_processed_dataset.py` re-checks all of it independently.

Split: the TEST set is the same 45 flakes as the fixed hold-out used in the report (results/splits/holdout_split.csv);
the remaining 133 are split 106 train / 27 val, stratified by class x data family x black (seed 42).
"""
import glob
import hashlib
import os
import shutil
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.model_selection import StratifiedShuffleSplit

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "Dataset"
OUT = RAW / "Processed"
SEED, VAL_N = 42, 27
GT_FIRST_ROW = 6   # data starts on Excel row 6 (header on row 5)


def read_sheet(path, id_col):
    """Return (header dict col_letter->text, rows list of dict(col_letter->value, '_row'->excel row)). Rows with empty id are skipped."""
    ws = load_workbook(path, data_only=True).worksheets[0]
    hdr = {c.column_letter: c.value for c in ws[5] if c.value is not None}
    rows = []
    for r in range(GT_FIRST_ROW, ws.max_row + 1):
        if ws[f"{id_col}{r}"].value in (None, ""):
            continue
        d = {c.column_letter: c.value for c in ws[r] if c.value is not None}
        d["_row"] = r
        rows.append(d)
    return hdr, rows


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    gt_hdr, gt_rows = read_sheet(RAW / "FTIR Classification_GT_Labels.xlsx", "C")
    col_hdr, col_rows = read_sheet(RAW / "Colorimeter_data.xlsx", "C")
    last_letter = max(gt_hdr, key=lambda k: (len(k), k))
    assert gt_hdr[last_letter].strip().lower().startswith("varified"), f"last GT column is {gt_hdr[last_letter]!r}"
    assert col_hdr["G"].startswith("Colorimenter L") and col_hdr["H"].startswith("Colorimenter a") and col_hdr["I"].startswith("Colorimenter b")
    col_by_id = {str(r["C"]).strip(): r for r in col_rows}
    assert len(col_by_id) == len(col_rows) and set(col_by_id) == {str(r["C"]).strip() for r in gt_rows}, "GT and colour sheets list different sample ids"

    # raw spectra files indexed by name (ignore anything already inside Processed)
    files = [p for p in glob.glob(str(RAW / "**" / "*"), recursive=True) if os.path.isfile(p) and "/Processed/" not in p]
    dpt = {os.path.basename(p)[:-4]: p for p in files if p.endswith(".dpt")}
    opus = {os.path.basename(p): p for p in files if not p.endswith((".dpt", ".xlsx", ".JPG", ".csv"))}

    recs = []
    for g in gt_rows:
        sid = str(g["C"]).strip()
        c = col_by_id[sid]
        raw_label = str(g[last_letter]).strip()
        key = raw_label.upper()
        label = {"PE": "PE", "PP": "PP"}.get(key)
        reason = ""
        if key == "PE-DOMINANT":
            reason = "ambiguous label: 'PE-dominant' (kept aside, not scored)"
        elif key == "UNCERTAIN":
            reason = "label 'Uncertain' (white Bopp shred; libraries say PVC / mineral)"
        elif label is None:
            reason = f"unknown label {raw_label!r}"
        d_key, o_key = sid + ".0", sid + ".0"
        if d_key not in dpt:
            reason = (reason + "; " if reason else "") + "no spectrum file"
        recs.append(dict(
            sample_id=sid, plate_code=str(g["B"]).strip(), sl_no=int(float(g["A"])), source=str(g["D"]).strip(), plate_location=str(g["E"]).strip(),
            visual_colour=str(g["F"]).strip().lower(), label=label or "", y_PP=(1 if label == "PP" else (0 if label == "PE" else "")), label_raw=raw_label,
            L=float(c["G"]), a=float(c["H"]), b=float(c["I"]), illuminant=str(c["J"]).strip(), observer_deg=int(float(c["K"])),
            polymerflow_id=str(g["J"]).strip(), polymerflow_score=float(g["I"]), openspecy_id=str(g["L"]).strip(), openspecy_score=float(g["K"]),
            gt_excel_row=g["_row"], colour_excel_row=c["_row"],
            original_dpt=os.path.relpath(dpt[d_key], RAW) if d_key in dpt else "", original_opus=os.path.relpath(opus[o_key], RAW) if o_key in opus else "",
            excluded_reason=reason))
    df = pd.DataFrame(recs)
    df["family"] = np.where(df.source.str.startswith("Impact"), "Impact", "Circular11")
    df["is_black"] = df.visual_colour.eq("black")

    # ---------------------------------------------------------------- splits
    use = df.excluded_reason.eq("")
    df["split"] = "excluded"
    ho = pd.read_csv(ROOT / "results" / "splits" / "holdout_split.csv").set_index("sample_id").split
    assert set(ho.index) == set(df.sample_id[use]), "hold-out file does not match the analysis set"
    df.loc[use, "split"] = df.loc[use, "sample_id"].map(ho).map({"test": "test", "dev": "dev"})
    dev = df[df.split == "dev"]
    strat = (dev.label + "|" + dev.family + "|" + np.where(dev.is_black, "blk", "col")).values
    tr_i, va_i = next(StratifiedShuffleSplit(1, test_size=VAL_N, random_state=SEED).split(dev, strat))
    df.loc[dev.index[tr_i], "split"] = "train"; df.loc[dev.index[va_i], "split"] = "val"
    assert set(df.split) == {"train", "val", "test", "excluded"}

    # ---------------------------------------------------------------- write
    if OUT.exists():
        shutil.rmtree(OUT)
    meta_cols = ["sample_id", "split", "label", "y_PP", "label_raw", "L", "a", "b", "visual_colour", "is_black", "source", "family", "plate_location",
                 "polymerflow_id", "polymerflow_score", "openspecy_id", "openspecy_score", "plate_code", "sl_no", "illuminant", "observer_deg",
                 "spectrum_file", "original_dpt", "original_opus", "gt_excel_row", "colour_excel_row", "excluded_reason"]
    df["spectrum_file"] = ""
    feats = {}
    for split in ["train", "val", "test"]:
        sub = df[df.split == split]
        (OUT / split / "spectra").mkdir(parents=True, exist_ok=True)
        rows, wn_ref = [], None
        for i, r in sub.iterrows():
            fn = f"{r.label}__{r.sample_id}.dpt"
            shutil.copy2(RAW / r.original_dpt, OUT / split / "spectra" / fn)
            df.at[i, "spectrum_file"] = f"{split}/spectra/{fn}"
            arr = np.loadtxt(RAW / r.original_dpt); o = np.argsort(arr[:, 0]); wn, ab = arr[o, 0], arr[o, 1]
            if wn_ref is None: wn_ref = wn
            assert np.allclose(wn, wn_ref, atol=1e-3), f"wavenumber grid differs for {r.sample_id}"
            rows.append(ab)
        m = df.loc[sub.index]
        X = pd.DataFrame(np.vstack(rows), columns=[f"wn_{w:.5f}" for w in wn_ref])
        head = pd.DataFrame({"sample_id": m.sample_id.values, "label": m.label.values, "y_PP": m.y_PP.values, "L": m.L.values, "a": m.a.values, "b": m.b.values})
        pd.concat([head, X], axis=1).to_csv(OUT / split / "features.csv", index=False)
        m[meta_cols].to_csv(OUT / split / "metadata.csv", index=False)
    # excluded + unlabelled extras
    (OUT / "excluded" / "spectra").mkdir(parents=True)
    ex = df[df.split == "excluded"]
    for i, r in ex.iterrows():
        if r.original_dpt:
            fn = f"{(r.label_raw.replace('/', '-') or 'NA')}__{r.sample_id}.dpt"
            shutil.copy2(RAW / r.original_dpt, OUT / "excluded" / "spectra" / fn); df.at[i, "spectrum_file"] = f"excluded/spectra/{fn}"
    df.loc[df.split == "excluded", meta_cols].to_csv(OUT / "excluded" / "metadata.csv", index=False)
    known = {s + ".0" for s in df.sample_id} | {s + ".0.dpt" for s in df.sample_id}
    extras = [p for p in files if os.path.basename(p) not in known and not p.endswith((".xlsx", ".JPG"))]
    (OUT / "excluded" / "unlabelled_extra_files").mkdir()
    ex_rows = []
    for p in sorted(extras):
        shutil.copy2(p, OUT / "excluded" / "unlabelled_extra_files" / os.path.basename(p))
        ex_rows.append(dict(file=os.path.basename(p), original=os.path.relpath(p, RAW), reason="no row in the GT sheet (re-scan / stray export); not used"))
    pd.DataFrame(ex_rows).to_csv(OUT / "excluded" / "unlabelled_extra_files.csv", index=False)

    df[meta_cols].to_csv(OUT / "all_metadata.csv", index=False)
    summ = (df.assign(n=1, label=df.label.replace("", "n/a")).pivot_table(index="split", columns="label", values="n", aggfunc="sum", fill_value=0)
            .join(df.groupby("split").agg(total=("sample_id", "size"), black=("is_black", "sum"))).reindex(["train", "val", "test", "excluded"]))
    fam = pd.crosstab(df.split, df.family).reindex(["train", "val", "test", "excluded"])
    summ = summ.join(fam); summ.to_csv(OUT / "split_summary.csv")
    print(summ.to_string()); print("excluded unlabelled extra files:", len(ex_rows))


if __name__ == "__main__":
    main()
