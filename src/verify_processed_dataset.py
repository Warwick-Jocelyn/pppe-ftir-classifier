"""Independent check of Dataset/Processed. Does NOT import the builder: it re-reads the raw Excel cells (by recorded row number),
the raw spectra and the native OPUS files, and compares them with what was written. Exit code 1 if anything fails.

Checks
 1  GT + colour values in metadata == the Excel cells at gt_excel_row / colour_excel_row (sample id, label_raw, colour, L a b, source, library IDs)
 2  label is the LAST GT column ("Varified Polymer ID"); PE/PP mapping; excluded rows really are PE-dominant / Uncertain / no spectrum
 3  every spectrum file in a split folder is byte-identical to its original .dpt, which is named <sample_id>.0.dpt
 4  the .dpt equals the native Bruker .0 file (brukeropus) -> the text export belongs to that instrument file
 5  features.csv rows == parsed spectra, in the same order as metadata.csv; L a b columns == metadata
 6  splits are disjoint, cover every analysis flake exactly once; test == the fixed hold-out used in the report
 7  filename prefix label == metadata label; no duplicate sample ids
 8  content cross-check: a simple band-ratio rule (720/1463 vs 1376/1463) computed from each spectrum agrees with its label,
    so shifted/mixed-up rows would show up as many disagreements
"""
import hashlib
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW, P = ROOT / "Dataset", ROOT / "Dataset" / "Processed"
fails, lines = [], []


def say(msg=""): lines.append(msg); print(msg)
def check(cond, msg):
    if not cond: fails.append(msg); say("  FAIL: " + msg)
    return cond


meta = pd.read_csv(P / "all_metadata.csv", keep_default_na=False)
gt = load_workbook(RAW / "FTIR Classification_GT_Labels.xlsx", data_only=True).worksheets[0]
cs = load_workbook(RAW / "Colorimeter_data.xlsx", data_only=True).worksheets[0]
hdr_last = [c for c in gt[5] if c.value is not None][-1]
say(f"GT sheet last column = {hdr_last.column_letter}: {hdr_last.value!r}")
check(str(hdr_last.value).strip().lower().startswith("varified"), "last GT column is not 'Varified Polymer ID'")
LAST = hdr_last.column_letter

say(f"\n[1,2] Excel cross-check of {len(meta)} records")
bad = 0
for _, r in meta.iterrows():
    g, c = int(r.gt_excel_row), int(r.colour_excel_row)
    ok = (str(gt[f"C{g}"].value).strip() == r.sample_id and str(cs[f"C{c}"].value).strip() == r.sample_id
          and str(gt[f"{LAST}{g}"].value).strip() == r.label_raw
          and np.isclose(float(cs[f"G{c}"].value), r.L) and np.isclose(float(cs[f"H{c}"].value), r.a) and np.isclose(float(cs[f"I{c}"].value), r.b)
          and str(gt[f"D{g}"].value).strip() == r.source and str(gt[f"E{g}"].value).strip() == str(r.plate_location)
          and str(gt[f"F{g}"].value).strip().lower() == r.visual_colour
          and str(gt[f"J{g}"].value).strip() == r.polymerflow_id and np.isclose(float(gt[f"I{g}"].value), r.polymerflow_score)
          and str(gt[f"L{g}"].value).strip() == r.openspecy_id and np.isclose(float(gt[f"K{g}"].value), r.openspecy_score))
    if not ok: bad += 1; check(False, f"{r.sample_id}: metadata differs from Excel rows {g}/{c}")
    lab = {"PE": "PE", "PP": "PP"}.get(r.label_raw.upper(), "")
    check(r.label == lab, f"{r.sample_id}: label {r.label!r} != mapping of {r.label_raw!r}")
    if r.split == "excluded":
        check(r.label_raw.upper() in ("PE-DOMINANT", "UNCERTAIN") or r.spectrum_file == "", f"{r.sample_id}: excluded without a valid reason")
say(f"  {len(meta)-bad}/{len(meta)} records identical to the Excel cells")
check(meta.sample_id.is_unique, "duplicate sample ids")

say("\n[3,4,5,7] spectra files")
from brukeropus import read_opus
n_ok = 0
for split in ["train", "val", "test"]:
    m = pd.read_csv(P / split / "metadata.csv", keep_default_na=False)
    f = pd.read_csv(P / split / "features.csv")
    check(list(f.sample_id) == list(m.sample_id), f"{split}: features.csv and metadata.csv are in a different order")
    check(np.allclose(f[["L", "a", "b"]].values, m[["L", "a", "b"]].values), f"{split}: L a b differ between features and metadata")
    check(list(f.label) == list(m.label), f"{split}: labels differ between features and metadata")
    wn = np.array([float(c[3:]) for c in f.columns if c.startswith("wn_")]); spec = f[[c for c in f.columns if c.startswith("wn_")]].values
    for i, r in m.iterrows():
        orig = RAW / r.original_dpt; out = P / r.spectrum_file
        ok = check(orig.name == r.sample_id + ".0.dpt", f"{r.sample_id}: original file name {orig.name}")
        ok &= check(out.name == f"{r.label}__{r.sample_id}.dpt" and out.parent.parent.name == split, f"{r.sample_id}: bad output path {out}")
        ok &= check(hashlib.sha256(orig.read_bytes()).digest() == hashlib.sha256(out.read_bytes()).digest(), f"{r.sample_id}: copied dpt differs from original")
        arr = np.loadtxt(out); o = np.argsort(arr[:, 0])
        ok &= check(np.allclose(arr[o, 1], spec[i], atol=1e-9) and np.allclose(arr[o, 0], wn, atol=1e-4), f"{r.sample_id}: features.csv row differs from its dpt")
        op = read_opus(str(RAW / r.original_opus)).a
        ok &= check(np.allclose(np.sort(np.asarray(op.x)), np.sort(arr[:, 0]), atol=1e-4) and np.abs(np.asarray(op.y)[np.argsort(np.asarray(op.x))] - arr[o, 1]).max() < 1e-4,
                    f"{r.sample_id}: dpt does not match native OPUS file")
        n_ok += bool(ok)
    say(f"  {split}: {len(m)} flakes, wavenumber {wn.min():.2f}-{wn.max():.2f} cm-1, {len(wn)} points")
say(f"  {n_ok} flakes fully consistent (byte-identical copy, equals native .0, features.csv row, naming)")

say("\n[6] splits")
S = {s: set(pd.read_csv(P / s / "metadata.csv", keep_default_na=False).sample_id) for s in ["train", "val", "test"]}
check(not (S["train"] & S["val"] or S["train"] & S["test"] or S["val"] & S["test"]), "splits overlap")
use = set(meta.sample_id[meta.split != "excluded"])
check(S["train"] | S["val"] | S["test"] == use and sum(len(v) for v in S.values()) == len(use), "splits do not cover the analysis set exactly once")
ho = pd.read_csv(ROOT / "results" / "splits" / "holdout_split.csv")
check(set(ho.sample_id[ho.split == "test"]) == S["test"], "test set differs from the fixed hold-out in results/splits")
check(set(ho.sample_id[ho.split == "dev"]) == S["train"] | S["val"], "train+val differs from hold-out dev set")
tab = meta.assign(label=meta.label.replace("", "n/a")).pivot_table(index="split", columns="label", values="sample_id", aggfunc="count", fill_value=0)
say(tab.to_string())

say("\n[8] content cross-check (label vs band-ratio rule computed from the spectrum itself)")
rows = []
for split in ["train", "val", "test"]:
    f = pd.read_csv(P / split / "features.csv"); wn = np.array([float(c[3:]) for c in f.columns if c.startswith("wn_")])
    X = f[[c for c in f.columns if c.startswith("wn_")]].values
    band = lambda c, hw=6: X[:, (wn > c - hw) & (wn < c + hw)].max(1)
    b = band(1463, 8)
    pred = np.where(band(1376) / b > band(720) / b, "PP", "PE")      # PP: strong 1376 CH3 band; PE: strong 720 rocking band
    rows.append(pd.DataFrame({"sample_id": f.sample_id, "label": f.label, "rule": pred}))
r = pd.concat(rows); dis = r[r.label != r.rule]
say(f"  rule agrees with label for {len(r)-len(dis)}/{len(r)} flakes; disagreements: {dis.sample_id.tolist()}")
check(len(dis) <= 3, f"too many content/label disagreements ({len(dis)}) - rows may be mismatched")

say("\nRESULT: " + ("ALL CHECKS PASSED" if not fails else f"{len(fails)} FAILURES"))
(P / "verification_report.txt").write_text("\n".join(lines) + "\n")
sys.exit(1 if fails else 0)
